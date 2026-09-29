from datetime import datetime, timedelta
import io
import shutil
from pathlib import Path
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.core.database import get_db
from app.core.config import settings
from app.core.dependencies import get_current_user, require_ce
from app.models.schema import RFQTender, TenderDocument, Application, RFQJob
from app.schemas.tender import (
    TenderCreate, TenderUpdate, TenderOut, TenderDocumentOut,
    RFQJobCreate, RFQJobUpdate, RFQJobOut
)

router = APIRouter(prefix="/api/tenders", tags=["Tenders"])

def build_tender_out(tender: RFQTender, db: Session) -> TenderOut:
    now = datetime.now()
    from_date = tender.quotation_from_date or now
    to_date = tender.quotation_to_date or (now + timedelta(days=365))
    is_open = (tender.status == "PUBLISHED") and (from_date <= now <= to_date)
    sub_count = db.query(Application).filter(Application.tender_id == tender.tender_id).count()
    docs = db.query(TenderDocument).filter(TenderDocument.tender_id == tender.tender_id).all()
    jobs = db.query(RFQJob).filter(RFQJob.tender_id == tender.tender_id).order_by(RFQJob.job_id.asc()).all()

    jobs_out = []
    for j in jobs:
        calc_amt = float(j.amount) if j.amount is not None else (float(j.estimated_cost) if j.estimated_cost is not None else None)
        calc_rate = float(j.unit_rate) if j.unit_rate is not None else (
            (calc_amt / float(j.estimated_quantity)) if (calc_amt and j.estimated_quantity and float(j.estimated_quantity) > 0) else None
        )
        jobs_out.append(
            RFQJobOut(
                job_id=j.job_id,
                tender_id=j.tender_id,
                job_code=j.job_code,
                job_name=j.job_name,
                job_description=j.job_description,
                category=j.category or "Supply Item",
                work_type=j.work_type,
                cl_number=j.cl_number,
                estimated_quantity=float(j.estimated_quantity) if j.estimated_quantity is not None else None,
                unit=j.unit or "NOS",
                unit_rate=calc_rate,
                amount=calc_amt,
                estimated_cost=calc_amt,
                completion_period=j.completion_period,
                status=j.status,
                created_at=j.created_at,
                updated_at=j.updated_at
            )
        )
    
    return TenderOut(
        tender_id=tender.tender_id,
        tender_ref_no=tender.tender_ref_no,
        title=tender.title,
        authority_name=tender.authority_name,
        background=tender.background,
        scope_of_work=tender.scope_of_work,
        total_elements=tender.total_elements,
        element_types=tender.element_types,
        max_weight_mt=float(tender.max_weight_mt) if tender.max_weight_mt is not None else None,
        min_weight_mt=float(tender.min_weight_mt) if tender.min_weight_mt is not None else None,
        avg_weight_mt=float(tender.avg_weight_mt) if tender.avg_weight_mt is not None else None,
        emd_amount=float(tender.emd_amount) if tender.emd_amount is not None else 0.0,
        completion_period=tender.completion_period,
        quotation_from_date=tender.quotation_from_date,
        quotation_to_date=tender.quotation_to_date,
        revealing_date=tender.revealing_date,
        quotation_valid_upto=tender.quotation_valid_upto or (to_date + timedelta(days=90)),
        validity_period=tender.validity_period or "90 Days from Quotation Opening",
        opening_date=tender.opening_date,
        contact_person=tender.contact_person,
        contact_email=tender.contact_email,
        contact_phone=tender.contact_phone,
        office_address=tender.office_address,
        status=tender.status,
        created_by=tender.created_by,
        created_at=tender.created_at,
        updated_at=tender.updated_at,
        documents=[TenderDocumentOut.from_orm(d) for d in docs],
        jobs=jobs_out,
        is_window_open=is_open,
        submission_count=sub_count
    )

@router.get("/template/excel")
def download_excel_template():
    """Generates and downloads the official 8-column RFQ Items Excel template."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RFQ Items Template"

    headers = [
        "Estimate Quantity (only Figures)",
        "Item Detailed Specification Description",
        "Work Type (eg. Earth Work, Electrical works.. etc - upto 200 Characters)",
        "Item Short Description (upto 100 Characters)",
        "APSS / Morth Cl. Number (upto 200 Characters)",
        "Rate (INR) (Upto 2 Decimals)",
        "UOM (upto 50 Characters)",
        "Amount (INR) (Upto 2 Decimals)"
    ]

    header_fill = PatternFill(start_color="7A1315", end_color="7A1315", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D0D7DE'),
        right=Side(style='thin', color='D0D7DE'),
        top=Side(style='thin', color='D0D7DE'),
        bottom=Side(style='thin', color='D0D7DE')
    )

    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Sample rows for user guidance
    sample_rows = [
        [
            100,
            "Providing and fixing powder coated mild steel shaft door with frame and shutter fabricated from skin-pass galvanized iron sheets conforming to IS:513 with PUF insulation.",
            "Civil & Architectural Finishes",
            "Mild Steel Shaft Door (GI Skin-Pass 1.20mm)",
            "APSS Cl. 1204 / MORTH 2000",
            10957.00,
            "Sqm",
            1095700.00
        ],
        [
            250,
            "Fabrication, supply, transportation and erection of structural steel members, trusses, columns, and rafters conforming to IS:2062 Grade E250.",
            "Structural Steel Works",
            "Structural Steel Fabrication & Erection",
            "APSS Cl. 1400 / IS:800",
            75000.00,
            "MT",
            18750000.00
        ]
    ]

    for row in sample_rows:
        ws.append(row)

    # Set column widths
    col_widths = [20, 45, 30, 30, 25, 18, 15, 20]
    for i, width in enumerate(col_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 40

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=RFQ_Items_Upload_Template.xlsx"}
    )

@router.get("", response_model=List[TenderOut])
def list_tenders(
    status_filter: Optional[str] = None,
    created_by: Optional[int] = None,
    db: Session = Depends(get_db)
):
    query = db.query(RFQTender)
    if status_filter and status_filter != "ALL":
        query = query.filter(RFQTender.status == status_filter)
    if created_by:
        query = query.filter(RFQTender.created_by == created_by)
    tenders = query.order_by(RFQTender.created_at.desc()).all()
    return [build_tender_out(t, db) for t in tenders]

@router.get("/{tender_id}", response_model=TenderOut)
def get_tender(tender_id: int, db: Session = Depends(get_db)):
    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")
    return build_tender_out(tender, db)

@router.post("", response_model=TenderOut)
def create_tender(
    payload: TenderCreate,
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    now = datetime.now()
    q_from = payload.quotation_from_date or now
    q_to = payload.quotation_to_date or (payload.opening_date or (now + timedelta(days=30)))
    if q_to <= q_from:
        q_to = q_from + timedelta(days=30)

    # Revealing date defaults to opening_date or quotation_to_date if not set
    rev_date = payload.revealing_date or payload.opening_date or q_to
    q_valid_upto = payload.quotation_valid_upto or (q_to + timedelta(days=90))
    valid_text = payload.validity_period or "90 Days from Quotation Opening"

    tender = RFQTender(
        tender_ref_no=payload.tender_ref_no.strip(),
        title=payload.title.strip(),
        authority_name=payload.authority_name.strip(),
        background=payload.background,
        scope_of_work=payload.scope_of_work,
        total_elements=payload.total_elements,
        element_types=payload.element_types,
        max_weight_mt=payload.max_weight_mt,
        min_weight_mt=payload.min_weight_mt,
        avg_weight_mt=payload.avg_weight_mt,
        emd_amount=payload.emd_amount if payload.emd_amount is not None else 0.0,
        completion_period=payload.completion_period,
        quotation_from_date=q_from,
        quotation_to_date=q_to,
        revealing_date=rev_date,
        quotation_valid_upto=q_valid_upto,
        validity_period=valid_text,
        opening_date=payload.opening_date or q_to,
        contact_person=payload.contact_person,
        contact_email=payload.contact_email,
        contact_phone=payload.contact_phone,
        office_address=payload.office_address,
        status=payload.status,
        created_by=current_user["id"]
    )
    db.add(tender)
    db.flush()

    # Add RFQ Items
    if payload.jobs:
        for idx, job_data in enumerate(payload.jobs):
            job_code = job_data.job_code or f"ITEM-{tender.tender_id}-{idx + 1}"
            
            qty = job_data.estimated_quantity
            unit_rate = job_data.unit_rate
            calc_amount = job_data.amount
            if calc_amount is None and qty and unit_rate:
                calc_amount = qty * unit_rate
            elif calc_amount is None:
                calc_amount = job_data.estimated_cost

            job_obj = RFQJob(
                tender_id=tender.tender_id,
                job_code=job_code,
                job_name=job_data.job_name,
                job_description=job_data.job_description,
                category=job_data.category or "Supply Item",
                work_type=job_data.work_type,
                cl_number=job_data.cl_number,
                estimated_quantity=qty,
                unit=job_data.unit or "NOS",
                unit_rate=unit_rate,
                amount=calc_amount,
                estimated_cost=calc_amount,
                completion_period=job_data.completion_period,
                status=job_data.status or "ACTIVE"
            )
            db.add(job_obj)

    db.commit()
    db.refresh(tender)
    return build_tender_out(tender, db)

@router.put("/{tender_id}", response_model=TenderOut)
def update_tender(
    tender_id: int,
    payload: TenderUpdate,
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")

    if current_user["role"] != "SUPER_ADMIN" and tender.created_by != current_user["id"]:
        raise HTTPException(status_code=403, detail="Access forbidden: Only the Officer who raised this RFQ can edit its details.")

    if payload.title is not None:
        tender.title = payload.title.strip()
    if payload.authority_name is not None:
        tender.authority_name = payload.authority_name.strip()
    if payload.background is not None:
        tender.background = payload.background
    if payload.scope_of_work is not None:
        tender.scope_of_work = payload.scope_of_work
    if payload.total_elements is not None:
        tender.total_elements = payload.total_elements
    if payload.element_types is not None:
        tender.element_types = payload.element_types
    if payload.max_weight_mt is not None:
        tender.max_weight_mt = payload.max_weight_mt
    if payload.min_weight_mt is not None:
        tender.min_weight_mt = payload.min_weight_mt
    if payload.avg_weight_mt is not None:
        tender.avg_weight_mt = payload.avg_weight_mt
    if payload.emd_amount is not None:
        tender.emd_amount = payload.emd_amount
    if payload.completion_period is not None:
        tender.completion_period = payload.completion_period
    if payload.quotation_from_date is not None:
        tender.quotation_from_date = payload.quotation_from_date
    if payload.quotation_to_date is not None:
        tender.quotation_to_date = payload.quotation_to_date
    if payload.revealing_date is not None:
        tender.revealing_date = payload.revealing_date
    if payload.quotation_valid_upto is not None:
        tender.quotation_valid_upto = payload.quotation_valid_upto
    if payload.validity_period is not None:
        tender.validity_period = payload.validity_period
    if payload.opening_date is not None:
        tender.opening_date = payload.opening_date
    if payload.contact_person is not None:
        tender.contact_person = payload.contact_person
    if payload.contact_email is not None:
        tender.contact_email = payload.contact_email
    if payload.contact_phone is not None:
        tender.contact_phone = payload.contact_phone
    if payload.office_address is not None:
        tender.office_address = payload.office_address
    if payload.status is not None:
        tender.status = payload.status

    tender.updated_at = datetime.now()

    if payload.jobs is not None:
        db.query(RFQJob).filter(RFQJob.tender_id == tender_id).delete()
        for idx, job_data in enumerate(payload.jobs):
            job_code = job_data.job_code or f"ITEM-{tender_id}-{idx + 1}"
            qty = job_data.estimated_quantity
            unit_rate = job_data.unit_rate
            calc_amount = job_data.amount
            if calc_amount is None and qty and unit_rate:
                calc_amount = qty * unit_rate
            elif calc_amount is None:
                calc_amount = job_data.estimated_cost

            job_obj = RFQJob(
                tender_id=tender_id,
                job_code=job_code,
                job_name=job_data.job_name,
                job_description=job_data.job_description,
                category=job_data.category or "Supply Item",
                work_type=job_data.work_type,
                cl_number=job_data.cl_number,
                estimated_quantity=qty,
                unit=job_data.unit or "NOS",
                unit_rate=unit_rate,
                amount=calc_amount,
                estimated_cost=calc_amount,
                completion_period=job_data.completion_period,
                status=job_data.status or "ACTIVE"
            )
            db.add(job_obj)

    db.commit()
    db.refresh(tender)
    return build_tender_out(tender, db)

@router.post("/{tender_id}/jobs", response_model=RFQJobOut)
def add_job_to_tender(
    tender_id: int,
    payload: RFQJobCreate,
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")

    if current_user["role"] != "SUPER_ADMIN" and tender.created_by != current_user["id"]:
        raise HTTPException(status_code=403, detail="Access forbidden: Only the Officer who raised this RFQ can add RFQ items.")

    job_count = db.query(RFQJob).filter(RFQJob.tender_id == tender_id).count()
    job_code = payload.job_code or f"ITEM-{tender_id}-{job_count + 1}"

    qty = payload.estimated_quantity
    unit_rate = payload.unit_rate
    calc_amount = payload.amount
    if calc_amount is None and qty and unit_rate:
        calc_amount = qty * unit_rate
    elif calc_amount is None:
        calc_amount = payload.estimated_cost

    job_obj = RFQJob(
        tender_id=tender_id,
        job_code=job_code,
        job_name=payload.job_name,
        job_description=payload.job_description,
        category=payload.category or "Supply Item",
        work_type=payload.work_type,
        cl_number=payload.cl_number,
        estimated_quantity=qty,
        unit=payload.unit or "NOS",
        unit_rate=unit_rate,
        amount=calc_amount,
        estimated_cost=calc_amount,
        completion_period=payload.completion_period,
        status=payload.status or "ACTIVE"
    )
    db.add(job_obj)
    db.commit()
    db.refresh(job_obj)
    return RFQJobOut.from_orm(job_obj)

@router.delete("/{tender_id}/jobs/{job_id}")
def delete_job(
    tender_id: int,
    job_id: int,
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")

    if current_user["role"] != "SUPER_ADMIN" and tender.created_by != current_user["id"]:
        raise HTTPException(status_code=403, detail="Access forbidden: Only the Officer who raised this RFQ can delete items.")

    job_obj = db.query(RFQJob).filter(RFQJob.job_id == job_id, RFQJob.tender_id == tender_id).first()
    if not job_obj:
        raise HTTPException(status_code=404, detail="RFQ Item not found")

    db.delete(job_obj)
    db.commit()
    return {"message": "RFQ item deleted successfully", "job_id": job_id}

@router.patch("/{tender_id}/status")
def update_tender_status(
    tender_id: int,
    new_status: str,
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    valid_statuses = ["DRAFT", "PUBLISHED", "CLOSED", "CANCELLED"]
    if new_status not in valid_statuses:
        raise HTTPException(status_code=400, detail=f"Invalid status. Must be one of: {valid_statuses}")

    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")

    if current_user["role"] != "SUPER_ADMIN" and tender.created_by != current_user["id"]:
        raise HTTPException(status_code=403, detail="Access forbidden: Only the Officer who raised this RFQ can change its status.")

    tender.status = new_status
    tender.updated_at = datetime.now()
    db.commit()
    return {"message": "Status updated successfully", "tender_id": tender_id, "status": new_status}

@router.post("/{tender_id}/documents")
def upload_tender_document(
    tender_id: int,
    document_type: str = Form("RFQ_DOCUMENT"), # RFQ_DOCUMENT | PAPER_CLIPPING
    file: UploadFile = File(...),
    current_user: dict = Depends(require_ce),
    db: Session = Depends(get_db)
):
    tender = db.query(RFQTender).filter(RFQTender.tender_id == tender_id).first()
    if not tender:
        raise HTTPException(status_code=404, detail="Quotation/RFQ not found")

    if current_user["role"] != "SUPER_ADMIN" and tender.created_by != current_user["id"]:
        raise HTTPException(status_code=403, detail="Access forbidden: Only the Officer who raised this RFQ can upload documents.")

    safe_filename = f"rfq_{tender_id}_{document_type.lower()}_{datetime.now().strftime('%Y%m%d%H%M%S')}_{file.filename}"
    file_dest = settings.UPLOAD_DIR / safe_filename

    with open(file_dest, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    file_size_kb = int(file_dest.stat().st_size / 1024)

    doc = TenderDocument(
        tender_id=tender_id,
        document_type=document_type,
        file_name=file.filename,
        file_path=f"/uploads/{safe_filename}",
        file_size_kb=file_size_kb,
        uploaded_by=current_user["id"]
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc
